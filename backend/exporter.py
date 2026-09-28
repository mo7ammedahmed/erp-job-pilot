import io
from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.lib.pagesizes import A4
from reportlab.lib.enums import TA_RIGHT, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import arabic_reshaper
from bidi.algorithm import get_display

pdfmetrics.registerFont(TTFont("Amiri", str(Path(__file__).parent / "fonts" / "Amiri-Regular.ttf")))
LABELS = {
    "en": {"summary": "Summary", "experience": "Experience", "education": "Education", "skills": "Skills", "projects": "Projects",
           "languages": "Languages", "certifications": "Certifications"},
    "ar": {"summary": "الملخص", "experience": "الخبرات", "education": "التعليم", "skills": "المهارات", "projects": "المشاريع",
           "languages": "اللغات", "certifications": "الشهادات"},
}


def blocks(cv, lang):
    L = LABELS[lang]
    c = cv.get("contact") or {}
    yield "title", cv.get("name", "")
    if cv.get("headline"):
        yield "sub", cv["headline"]
    contact = " | ".join(x for x in [c.get("email"), c.get("phone"), c.get("location"), *(c.get("links") or [])] if x)
    if contact:
        yield "sub", contact
    if cv.get("summary"):
        yield "h", L["summary"]
        yield "p", cv["summary"]
    if cv.get("experience"):
        yield "h", L["experience"]
        for e in cv["experience"]:
            if not e:
                continue
            yield "b", " — ".join(x for x in [e.get("title"), e.get("company")] if x)
            meta = " | ".join(x for x in [f"{e.get('start', '')} – {e.get('end', '')}".strip(" –"), e.get("location")] if x)
            if meta:
                yield "sub", meta
            for bl in e.get("bullets") or []:
                if bl:
                    yield "li", bl
    if cv.get("education"):
        yield "h", L["education"]
        for ed in cv["education"]:
            if ed:
                yield "p", " — ".join(x for x in [" ".join(filter(None, [ed.get("degree"), ed.get("field")])), ed.get("institution"),
                                                  f"{ed.get('start', '')} – {ed.get('end', '')}".strip(" –")] if x)
    if cv.get("skills"):
        yield "h", L["skills"]
        yield "p", " • ".join(s for s in cv["skills"] if s)
    if cv.get("projects"):
        yield "h", L["projects"]
        for p in cv["projects"]:
            if p:
                yield "li", ": ".join(x for x in [p.get("name"), p.get("description")] if x)
    if cv.get("languages"):
        yield "h", L["languages"]
        yield "p", " • ".join(" ".join(filter(None, [l.get("name"), f"({l['level']})" if l.get("level") else ""])) for l in cv["languages"] if l)
    if cv.get("certifications"):
        yield "h", L["certifications"]
        for ce in cv["certifications"]:
            if ce:
                yield "li", " — ".join(x for x in [ce.get("name"), ce.get("issuer"), str(ce.get("year") or "")] if x)


def _items(cv, cover, lang, doc):
    if doc == "cover":
        return [("title", cv.get("name", ""))] + [("p", para) for para in (cover or "").split("\n") if para.strip()]
    return list(blocks(cv, lang))


def _rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rpr.append(OxmlElement("w:rtl"))


def to_docx(cv, cover, lang, doc="cv"):
    d = Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10.5)
    for kind, text in _items(cv, cover, lang, doc):
        if kind == "title":
            p = d.add_heading(text, level=0)
        elif kind == "h":
            p = d.add_heading(text, level=2)
        elif kind == "li":
            p = d.add_paragraph(text, style="List Bullet")
        else:
            p = d.add_paragraph()
            r = p.add_run(text)
            r.bold = kind == "b"
            r.italic = kind == "sub"
        if lang == "ar":
            _rtl(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pdf(cv, cover, lang, doc="cv"):
    ar = lang == "ar"
    font = "Amiri" if ar else "Helvetica"
    bold = "Amiri" if ar else "Helvetica-Bold"
    align = TA_RIGHT if ar else TA_LEFT
    st = {
        "title": ParagraphStyle("t", fontName=bold, fontSize=18, leading=24, alignment=align, textColor="#064E3B"),
        "sub": ParagraphStyle("s", fontName=font, fontSize=9.5, leading=13, alignment=align, textColor="#475569"),
        "h": ParagraphStyle("h", fontName=bold, fontSize=12.5, leading=18, alignment=align, spaceBefore=10, textColor="#064E3B"),
        "b": ParagraphStyle("b", fontName=bold, fontSize=10.5, leading=15, alignment=align, spaceBefore=6),
        "p": ParagraphStyle("p", fontName=font, fontSize=10, leading=15, alignment=align),
        "li": ParagraphStyle("l", fontName=font, fontSize=10, leading=15, alignment=align, leftIndent=0 if ar else 10, rightIndent=10 if ar else 0),
    }
    story = []
    for kind, text in _items(cv, cover, lang, doc):
        t = f"• {text}" if kind == "li" and not ar else (f"{text} •" if kind == "li" else text)
        if ar:
            t = get_display(arabic_reshaper.reshape(t))
        story.append(Paragraph(escape(t), st[kind]))
        if doc == "cover" and kind == "p":
            story.append(Spacer(1, 6))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=40, rightMargin=40, topMargin=36, bottomMargin=36).build(story)
    return buf.getvalue()
